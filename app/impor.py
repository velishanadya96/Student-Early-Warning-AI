"""Impor CSV / Excel (FR-02, FR-03). Semua-atau-tidak-sama-sekali: bila ada baris salah, tidak ada yang disimpan
dan semua kesalahan dikembalikan (nomor baris + kolom + pesan) agar bisa diperbaiki sekaligus.

Format akademik  (satu baris = satu mahasiswa x satu mata kuliah, per periode):
  wajib   : nim, kode_mk, nilai_mk, total_pertemuan, hadir
  opsional: nama, prodi, angkatan, ipk, nidn_dosen_pa      (nama wajib bila NIM belum ada di sistem)
Format tugas     (satu baris = satu tugas):
  wajib   : nim, kode_mk
  opsional: judul, waktu_kumpul (kosong = belum dikumpulkan), status_terlambat (tepat / terlambat)"""
import csv
import io
from datetime import datetime

from .errors import ApiError
from .extensions import db
from .models import DosenPA, Mahasiswa, Nilai, Presensi, Tugas

MAKS_BARIS = 20000
MAKS_ERROR = 50
KOLOM = {
    "akademik": {"wajib": ["nim", "kode_mk", "nilai_mk", "total_pertemuan", "hadir"],
                 "opsional": ["nama", "prodi", "angkatan", "ipk", "nidn_dosen_pa"]},
    "tugas": {"wajib": ["nim", "kode_mk"], "opsional": ["judul", "waktu_kumpul", "status_terlambat"]},
}
ALIAS = {"kode mk": "kode_mk", "kodemk": "kode_mk", "nilai": "nilai_mk", "pertemuan": "total_pertemuan",
         "total": "total_pertemuan", "kehadiran": "hadir", "gpa": "ipk", "nidn": "nidn_dosen_pa", "dosen_pa": "nidn_dosen_pa",
         "tanggal_kumpul": "waktu_kumpul", "dikumpulkan": "waktu_kumpul"}


def _norm_header(h):
    h = (h or "").strip().lower().replace("﻿", "").replace("-", "_")
    h = " ".join(h.split())
    h = ALIAS.get(h, h).replace(" ", "_")
    return ALIAS.get(h, h)


def baca_file(file_storage):
    """Baca .csv / .xlsx menjadi list of dict (kunci = header ternormalisasi) + nomor baris asli (header = baris 1)."""
    nama = (file_storage.filename or "").lower()
    data = file_storage.read()
    if not data:
        raise ApiError("FILE_KOSONG", "File kosong.", 422)
    if nama.endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            ws = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
            baris = [list(r) for r in ws.iter_rows(values_only=True)]
        except Exception:
            raise ApiError("FILE_TIDAK_VALID", "File Excel tidak bisa dibaca. Pastikan formatnya .xlsx.", 422)
    elif nama.endswith(".csv") or nama.endswith(".txt"):
        try:
            teks = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            teks = data.decode("latin-1")
        sample = teks[:2048]
        delim = ";" if sample.count(";") > sample.count(",") else ("\t" if sample.count("\t") > sample.count(",") else ",")
        baris = list(csv.reader(io.StringIO(teks), delimiter=delim))
    else:
        raise ApiError("FORMAT_TIDAK_DIDUKUNG", "Gunakan file .csv atau .xlsx.", 415)
    baris = [r for r in baris if any(str(c).strip() for c in r if c is not None)]
    if len(baris) < 2:
        raise ApiError("FILE_KOSONG", "File tidak berisi data (hanya header atau kosong).", 422)
    if len(baris) - 1 > MAKS_BARIS:
        raise ApiError("FILE_TERLALU_BESAR", f"Maksimal {MAKS_BARIS} baris per impor.", 413)
    header = [_norm_header(str(h) if h is not None else "") for h in baris[0]]
    hasil = []
    for i, r in enumerate(baris[1:], start=2):
        d = {}
        for j, h in enumerate(header):
            if h and j < len(r):
                v = r[j]
                d[h] = v.strip() if isinstance(v, str) else v
        hasil.append((i, d))
    return header, hasil


def _kosong(v):
    return v is None or (isinstance(v, str) and not v.strip())


def _int(v, lo, hi):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    n = int(str(v).strip())
    if not lo <= n <= hi:
        raise ValueError
    return n


def _float(v, lo, hi):
    n = float(str(v).strip().replace(",", "."))
    if not lo <= n <= hi:
        raise ValueError
    return n


def _tanggal(v):
    if isinstance(v, datetime):
        return v
    s = str(v).strip()
    for f in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d/%m/%Y %H:%M"):
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    raise ValueError


def _nim(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    if not s or len(s) > 20 or not s.isalnum():
        raise ValueError
    return s


def _cek_header(header, jenis):
    kurang = [k for k in KOLOM[jenis]["wajib"] if k not in header]
    if kurang:
        raise ApiError("KOLOM_KURANG", "Kolom wajib tidak ditemukan: " + ", ".join(kurang) + ".", 422,
                       {"kolom_wajib": KOLOM[jenis]["wajib"], "kolom_opsional": KOLOM[jenis]["opsional"],
                        "kolom_terbaca": [h for h in header if h]})


def _gagal(errors, total):
    raise ApiError("IMPOR_GAGAL", f"{len(errors)} baris bermasalah. Tidak ada data yang disimpan; perbaiki lalu unggah ulang.",
                   422, {"errors": errors[:MAKS_ERROR], "total_error": len(errors), "total_baris": total})


def impor_akademik(file_storage, periode):
    header, rows = baca_file(file_storage)
    _cek_header(header, "akademik")
    errors, valid = [], []
    dosen_by_nidn = {d.nidn: d for d in DosenPA.query.all()}
    mhs_ada = {m.nim for m in Mahasiswa.query.all()}
    kunci_terlihat = set()
    for no, r in rows:
        def err(kolom, pesan):
            errors.append({"baris": no, "kolom": kolom, "pesan": pesan})
        n0 = len(errors)
        try:
            nim = _nim(r.get("nim"))
        except (ValueError, TypeError):
            err("nim", "NIM wajib diisi (huruf/angka, maks. 20 karakter).")
            nim = None
        kode = str(r.get("kode_mk") or "").strip()
        if not kode or len(kode) > 20:
            err("kode_mk", "Kode mata kuliah wajib diisi (maks. 20 karakter).")
        try:
            nilai = _float(r.get("nilai_mk"), 0, 100)
        except (ValueError, TypeError):
            err("nilai_mk", "Nilai harus angka 0-100.")
            nilai = None
        try:
            total = _int(r.get("total_pertemuan"), 1, 60)
        except (ValueError, TypeError):
            err("total_pertemuan", "Total pertemuan harus bilangan bulat 1-60.")
            total = None
        try:
            hadir = _int(r.get("hadir"), 0, 60)
            if total is not None and hadir > total:
                err("hadir", f"Jumlah hadir ({hadir}) melebihi total pertemuan ({total}).")
        except (ValueError, TypeError):
            err("hadir", "Jumlah hadir harus bilangan bulat 0-60.")
            hadir = None
        ipk = angkatan = None
        if not _kosong(r.get("ipk")):
            try:
                ipk = _float(r.get("ipk"), 0, 4)
            except (ValueError, TypeError):
                err("ipk", "IPK harus angka 0-4 (pakai titik, mis. 3.25).")
        if not _kosong(r.get("angkatan")):
            try:
                angkatan = _int(r.get("angkatan"), 2000, 2100)
            except (ValueError, TypeError):
                err("angkatan", "Angkatan harus tahun 4 digit, mis. 2024.")
        nidn = None
        if not _kosong(r.get("nidn_dosen_pa")):
            nidn = str(r.get("nidn_dosen_pa")).strip()
            if isinstance(r.get("nidn_dosen_pa"), float):
                nidn = str(int(r["nidn_dosen_pa"]))
            if nidn not in dosen_by_nidn and nidn.zfill(10) in dosen_by_nidn:
                nidn = nidn.zfill(10)  # Excel membuang angka 0 di depan NIDN
            if nidn not in dosen_by_nidn:
                err("nidn_dosen_pa", f"NIDN {nidn} tidak terdaftar sebagai Dosen PA.")
        nama = str(r.get("nama") or "").strip() or None
        if nim and nim not in mhs_ada and not nama:
            err("nama", f"NIM {nim} belum ada di sistem; kolom nama wajib diisi.")
        if nim and kode and (nim, kode) in kunci_terlihat:
            err("kode_mk", f"Pasangan NIM {nim} + mata kuliah {kode} muncul dua kali.")
        kunci_terlihat.add((nim, kode))
        if len(errors) == n0:
            valid.append(dict(nim=nim, kode=kode, nilai=nilai, total=total, hadir=hadir, ipk=ipk, angkatan=angkatan,
                              nama=nama, prodi=(str(r.get("prodi") or "").strip() or None), nidn=nidn))
    if errors:
        _gagal(errors, len(rows))

    dibuat = {}
    for v in valid:
        m = dibuat.get(v["nim"]) or db.session.get(Mahasiswa, v["nim"])
        if not m:
            m = Mahasiswa(nim=v["nim"], nama=v["nama"])
            db.session.add(m)
        dibuat[v["nim"]] = m
        if v["nama"] and not m.id_user:  # nama SIAKAD hanya menimpa bila belum punya akun
            m.nama = v["nama"]
        if v["prodi"]:
            m.prodi = v["prodi"]
        if v["angkatan"]:
            m.angkatan = v["angkatan"]
        if v["ipk"] is not None:
            m.ipk = v["ipk"]
        if v["nidn"]:
            m.id_dosen_pa = dosen_by_nidn[v["nidn"]].id_dosen_pa
    db.session.flush()
    nims = {v["nim"] for v in valid}
    # Impor ulang periode yang sama menggantikan baris lama untuk pasangan (nim, mata kuliah) yang sama.
    for v in valid:
        Nilai.query.filter_by(nim=v["nim"], kode_mk=v["kode"], semester=periode).delete()
        Presensi.query.filter_by(nim=v["nim"], kode_mk=v["kode"], periode=periode).delete()
        db.session.add(Nilai(nim=v["nim"], kode_mk=v["kode"], semester=periode, nilai_mk=v["nilai"]))
        db.session.add(Presensi(nim=v["nim"], kode_mk=v["kode"], periode=periode,
                                total_pertemuan=v["total"], hadir=v["hadir"]))
    db.session.commit()
    mhs_baru = len(nims - mhs_ada)
    return {"periode": periode, "baris": len(valid), "mahasiswa": len(nims), "mahasiswa_baru": mhs_baru}


def impor_tugas(file_storage, periode):
    header, rows = baca_file(file_storage)
    _cek_header(header, "tugas")
    errors, valid = [], []
    mhs_ada = {m.nim for m in Mahasiswa.query.all()}
    for no, r in rows:
        n0 = len(errors)

        def err(kolom, pesan):
            errors.append({"baris": no, "kolom": kolom, "pesan": pesan})
        try:
            nim = _nim(r.get("nim"))
            if nim not in mhs_ada:
                err("nim", f"NIM {nim} belum ada di sistem. Impor data akademik dulu.")
        except (ValueError, TypeError):
            err("nim", "NIM wajib diisi (huruf/angka, maks. 20 karakter).")
            nim = None
        kode = str(r.get("kode_mk") or "").strip()
        if not kode or len(kode) > 20:
            err("kode_mk", "Kode mata kuliah wajib diisi (maks. 20 karakter).")
        kumpul = None
        if not _kosong(r.get("waktu_kumpul")):
            try:
                kumpul = _tanggal(r.get("waktu_kumpul"))
            except (ValueError, TypeError):
                err("waktu_kumpul", "Format tanggal: 2026-09-30 atau 2026-09-30 14:30 (kosongkan bila belum dikumpulkan).")
        status = str(r.get("status_terlambat") or "").strip().lower() or None
        if status and status not in ("tepat", "terlambat"):
            err("status_terlambat", "Isi 'tepat' atau 'terlambat' (atau kosongkan).")
        judul = (str(r.get("judul") or "").strip() or None)
        if judul and len(judul) > 150:
            err("judul", "Judul maksimal 150 karakter.")
        if len(errors) == n0:
            valid.append(dict(nim=nim, kode=kode, kumpul=kumpul, status=status, judul=judul))
    if errors:
        _gagal(errors, len(rows))
    nims = {v["nim"] for v in valid}
    # Impor ulang mengganti seluruh tugas periode ini untuk mahasiswa yang ada di file.
    for nim in nims:
        Tugas.query.filter_by(nim=nim, periode=periode).delete()
    for v in valid:
        db.session.add(Tugas(nim=v["nim"], kode_mk=v["kode"], periode=periode, judul=v["judul"],
                             waktu_kumpul=v["kumpul"], status_terlambat=v["status"]))
    db.session.commit()
    return {"periode": periode, "baris": len(valid), "mahasiswa": len(nims),
            "terkumpul": sum(1 for v in valid if v["kumpul"])}
