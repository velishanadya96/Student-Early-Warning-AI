import io
from datetime import date

from flask import request, send_file

from . import laporan_bp
from .common import periode_valid, skor_terbaru, urut_risiko
from .. import risk
from ..access import current_user, mahasiswa_query
from ..errors import ApiError
from ..models import Mahasiswa
from ..security import roles_required

KOLOM = ["NIM", "Nama", "Angkatan", "Dosen PA", "IPK", "Rata-rata nilai", "Kehadiran (%)", "Tugas terkumpul (%)",
         "Skor risiko", "Kategori", "Penyebab"]


def _baris(user, periode):
    mhs = mahasiswa_query(user).order_by(Mahasiswa.nim).all()
    skor = skor_terbaru([m.nim for m in mhs], periode)
    out = []
    for m in mhs:
        rs = skor.get(m.nim)
        if periode and not rs:
            continue
        f = risk.skor_json(rs)["fitur"] if rs else {}
        pct = lambda v: None if v is None else round(v * 100, 1)
        out.append({"nim": m.nim, "kategori": rs.kategori_risiko if rs else None, "skor": rs.skor if rs else None,
                    "sel": [m.nim, m.nama, m.angkatan, m.dosen_pa.nama if m.dosen_pa else "-", f.get("GPA"),
                            None if f.get("avg_course_grade") is None else round(f["avg_course_grade"], 1),
                            pct(f.get("attendance_rate")), pct(f.get("assignment_submission_rate")),
                            None if not rs else round(rs.skor, 3), rs.kategori_risiko if rs else "Belum dihitung",
                            (rs.faktor_dominan or "-") if rs else "-"]})
    return urut_risiko(out)


@laporan_bp.get("")
@roles_required("kaprodi", "admin", "dosen")
def laporan():
    """Laporan risiko mahasiswa (FR-09). ?periode=2026-1&format=xlsx|pdf. Dosen hanya memuat mahasiswa bimbingannya."""
    user = current_user()
    fmt = (request.args.get("format") or "xlsx").lower()
    if fmt not in ("xlsx", "pdf"):
        raise ApiError("VALIDASI_GAGAL", "Format harus xlsx atau pdf.", 422)
    periode = periode_valid(request.args.get("periode")) if request.args.get("periode") else (risk.periode_tersedia() or [None])[0]
    if not periode:
        raise ApiError("BELUM_ADA_DATA", "Belum ada data untuk dilaporkan.", 409)
    baris = _baris(user, periode)
    nama = f"laporan-risiko-{periode}.{fmt}"
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = f"Risiko {periode}"
        ws.append(KOLOM)
        for c in ws[1]:
            c.font = Font(bold=True)
        warna = {"High": "F8C9C9", "Medium": "FFE9B3", "Low": "CDEBD2"}
        for b in baris:
            ws.append(b["sel"])
            if b["kategori"] in warna:
                ws.cell(ws.max_row, 10).fill = PatternFill("solid", fgColor=warna[b["kategori"]])
        for i, lebar in enumerate([12, 26, 9, 22, 7, 14, 13, 18, 11, 11, 44], 1):
            ws.column_dimensions[chr(64 + i)].width = lebar
        buf = io.BytesIO()
        wb.save(buf)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
        st = getSampleStyleSheet()
        ringkas = {k: sum(1 for b in baris if b["kategori"] == k) for k in ("High", "Medium", "Low")}
        data = [KOLOM] + [[("-" if v is None else str(v)) for v in b["sel"]] for b in baris]
        tbl = Table(data, repeatRows=1)
        gaya = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey)]
        for i, b in enumerate(baris, 1):
            if b["kategori"] == "High":
                gaya.append(("BACKGROUND", (9, i), (9, i), colors.HexColor("#f8c9c9")))
        tbl.setStyle(TableStyle(gaya))
        doc.build([Paragraph(f"Laporan Risiko Mahasiswa - Periode {periode}", st["Title"]),
                   Paragraph(f"Dibuat {date.today().isoformat()} | Tinggi: {ringkas['High']} | Sedang: {ringkas['Medium']} | Rendah: {ringkas['Low']}", st["Normal"]),
                   Spacer(1, 10), tbl])
        mime = "application/pdf"
    buf.seek(0)
    return send_file(buf, mimetype=mime, as_attachment=True, download_name=nama)
