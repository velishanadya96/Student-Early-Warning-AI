"""Impor CSV -> hitung risiko -> notifikasi WA ke orang tua (dengan alasan) -> dashboard, catatan, laporan."""
import io

import pytest

from app import create_app
from app.extensions import db
from app.models import DosenPA, Mahasiswa, Notifikasi, OrangTua, RiskScore, User
from app.security import hash_password

AKADEMIK = (
    "nim,nama,prodi,angkatan,ipk,nidn_dosen_pa,kode_mk,nilai_mk,total_pertemuan,hadir\n"
    # Alya: baik semua -> Low
    "2024001,Alya,D3 TI,2024,3.7,0012345601,TI101,88,14,14\n"
    "2024001,Alya,D3 TI,2024,3.7,0012345601,TI102,86,14,13\n"
    # Citra: hanya kehadiran buruk -> High karena kehadiran
    "2024003,Citra,D3 TI,2024,3.1,0012345601,TI101,76,14,7\n"
    "2024003,Citra,D3 TI,2024,3.1,0012345601,TI102,78,14,8\n"
    # Eka: hanya IPK buruk, hadir bagus -> High bukan karena kehadiran
    "2024005,Eka,D3 TI,2023,2.1,0012345601,TI101,60,14,14\n"
    "2024005,Eka,D3 TI,2023,2.1,0012345601,TI102,62,14,13\n"
    # Zaki: milik dosen lain
    "2024009,Zaki,D3 TI,2023,3.5,0012345602,TI101,85,14,14\n"
)
TUGAS = (
    "nim,kode_mk,judul,waktu_kumpul,status_terlambat\n"
    "2024001,TI101,T1,2026-09-10 10:00,tepat\n2024001,TI102,T2,2026-09-11,tepat\n"
    "2024003,TI101,T1,2026-09-10,tepat\n2024003,TI102,T2,2026-09-12,terlambat\n"
    "2024005,TI101,T1,2026-09-10,tepat\n2024005,TI102,T2,2026-09-11,tepat\n"
    "2024009,TI101,T1,2026-09-10,tepat\n"
)


@pytest.fixture()
def app():
    app = create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite:///:memory:", "RATELIMIT_ENABLED": False,
                      "EMAIL_VERIFICATION_REQUIRED": False, "WA_PROVIDER": "manual"})
    with app.app_context():
        db.create_all()

        def mk(username, email, role):
            u = User(username=username, email=email, role=role, password_hash=hash_password("rahasia123"))
            db.session.add(u); db.session.flush(); return u
        mk("admin", "admin@x.id", "admin")
        mk("kaprodi", "kaprodi@x.id", "kaprodi")
        d1 = mk("dosen", "dosen@x.id", "dosen")
        d2 = mk("dosen2", "dosen2@x.id", "dosen")
        db.session.add_all([DosenPA(id_user=d1.id_user, nidn="0012345601", nama="Dosen Satu", no_wa="6281200000001"),
                            DosenPA(id_user=d2.id_user, nidn="0012345602", nama="Dosen Dua", no_wa="6281200000009")])
        # Citra sudah punya akun mahasiswa + orang tua dengan nomor WA; Eka punya orang tua tanpa nomor
        m = mk("citra", "citra@x.id", "mahasiswa")
        db.session.add(Mahasiswa(nim="2024003", id_user=m.id_user, nama="Citra Lestari"))
        db.session.flush()
        for nim, user, nama, wa in (("2024003", "ortu-citra", "Ibu Citra", "6281200000003"),):
            o = mk(user, f"{user}@x.id", "orangtua")
            db.session.add(OrangTua(id_user=o.id_user, nim=nim, nama=nama, no_wa=wa))
        db.session.add(Mahasiswa(nim="2024005", nama="Eka"))
        db.session.flush()
        o = mk("ortu-eka", "ortu-eka@x.id", "orangtua")
        db.session.add(OrangTua(id_user=o.id_user, nim="2024005", nama="Ibu Eka", no_wa=None))
        db.session.commit()
    return app


@pytest.fixture()
def c(app):
    return app.test_client()


def tok(c, email, **kw):
    r = c.post("/api/v1/auth/login", json={"email": email, "password": "rahasia123"})
    assert r.status_code == 200, r.get_json()
    return {"Authorization": "Bearer " + r.get_json()["access_token"]}


def upload(c, h, path, isi, periode="2026-1", nama="data.csv"):
    return c.post(path, headers=h, data={"periode": periode, "file": (io.BytesIO(isi.encode()), nama)},
                  content_type="multipart/form-data")


@pytest.fixture()
def terisi(c):
    h = tok(c, "admin@x.id")
    assert upload(c, h, "/api/v1/import/akademik", AKADEMIK).status_code == 201
    assert upload(c, h, "/api/v1/import/tugas", TUGAS).status_code == 201
    r = c.post("/api/v1/risk/hitung", headers=h, json={"periode": "2026-1"})
    assert r.status_code == 200
    return r.get_json()


# ---------------------------------------------------------------- impor
def test_impor_membuat_mahasiswa_baru_dan_memakai_nidn(c, app):
    h = tok(c, "admin@x.id")
    r = upload(c, h, "/api/v1/import/akademik", AKADEMIK)
    assert r.status_code == 201
    b = r.get_json()
    assert b["mahasiswa"] == 4 and b["mahasiswa_baru"] == 2 and b["baris"] == 7
    with app.app_context():
        assert db.session.get(Mahasiswa, "2024001").id_dosen_pa is not None
        assert db.session.get(Mahasiswa, "2024003").nama == "Citra Lestari"  # nama akun tidak ditimpa CSV


def test_impor_ulang_tidak_menggandakan(c, app):
    h = tok(c, "admin@x.id")
    upload(c, h, "/api/v1/import/akademik", AKADEMIK)
    upload(c, h, "/api/v1/import/akademik", AKADEMIK)
    from app.models import Nilai, Presensi
    with app.app_context():
        assert Nilai.query.count() == 7 and Presensi.query.count() == 7


@pytest.mark.parametrize("baris,kolom", [
    ("2024001,Alya,D3,2024,3.7,0012345601,TI101,150,14,14", "nilai_mk"),
    ("2024001,Alya,D3,2024,3.7,0012345601,TI101,80,14,20", "hadir"),
    ("2024001,Alya,D3,2024,5.2,0012345601,TI101,80,14,10", "ipk"),
    ("2024001,Alya,D3,2024,3.7,9999999999,TI101,80,14,10", "nidn_dosen_pa"),
    ("2024777,,D3,2024,3.7,0012345601,TI101,80,14,10", "nama"),
    ("abc def,Alya,D3,2024,3.7,0012345601,TI101,80,14,10", "nim"),
])
def test_impor_baris_salah_ditolak_semua(c, app, baris, kolom):
    h = tok(c, "admin@x.id")
    isi = "nim,nama,prodi,angkatan,ipk,nidn_dosen_pa,kode_mk,nilai_mk,total_pertemuan,hadir\n" + \
          "2024002,Bagas,D3,2024,3.0,0012345601,TI101,80,14,14\n" + baris + "\n"
    r = upload(c, h, "/api/v1/import/akademik", isi)
    assert r.status_code == 422
    b = r.get_json()
    assert b["code"] == "IMPOR_GAGAL"
    assert b["errors"][0]["baris"] == 3 and b["errors"][0]["kolom"] == kolom
    with app.app_context():
        assert Mahasiswa.query.filter_by(nim="2024002").first() is None  # atomik: baris benar pun tidak tersimpan


def test_impor_kolom_kurang_dan_format(c):
    h = tok(c, "admin@x.id")
    r = upload(c, h, "/api/v1/import/akademik", "nim,nama\n1,A\n")
    assert r.status_code == 422 and r.get_json()["code"] == "KOLOM_KURANG"
    assert upload(c, h, "/api/v1/import/akademik", "x", nama="data.pdf").status_code == 415
    assert upload(c, h, "/api/v1/import/akademik", AKADEMIK, periode="2026").status_code == 422


def test_impor_titik_koma_dan_xlsx(c):
    h = tok(c, "admin@x.id")
    assert upload(c, h, "/api/v1/import/akademik", AKADEMIK.replace(",", ";").replace("D3 TI", "D3 TI")).status_code == 201
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    for line in AKADEMIK.strip().split("\n"):
        ws.append([int(x) if x.isdigit() and not x.startswith("0") else x for x in line.split(",")])
    buf = io.BytesIO(); wb.save(buf)
    r = c.post("/api/v1/import/akademik", headers=h, content_type="multipart/form-data",
               data={"periode": "2026-2", "file": (io.BytesIO(buf.getvalue()), "d.xlsx")})
    assert r.status_code == 201, r.get_json()


def test_impor_tugas_nim_tak_dikenal(c):
    h = tok(c, "admin@x.id")
    r = upload(c, h, "/api/v1/import/tugas", "nim,kode_mk,judul,waktu_kumpul,status_terlambat\n999,TI1,x,,\n")
    assert r.status_code == 422 and r.get_json()["errors"][0]["kolom"] == "nim"


def test_impor_hanya_admin(c):
    for email in ("dosen@x.id", "kaprodi@x.id", "citra@x.id"):
        assert upload(c, tok(c, email), "/api/v1/import/akademik", AKADEMIK).status_code == 403
    assert c.post("/api/v1/import/akademik").status_code == 401


# ---------------------------------------------------------------- risiko
def test_hitung_kategori_dan_penyebab(c, terisi):
    assert terisi["Low"] == 2 and terisi["High"] == 2
    h = tok(c, "dosen@x.id")
    d = {m["nim"]: m for m in c.get("/api/v1/dashboard/dosen", headers=h).get_json()["mahasiswa"]}
    assert d["2024001"]["kategori"] == "Low" and d["2024001"]["faktor"] == []
    # Citra: penyebab = kehadiran saja
    assert [f["kode"] for f in d["2024003"]["faktor"]] == ["attendance_rate"]
    assert d["2024003"]["fitur"]["attendance_rate"] == pytest.approx(15 / 28, abs=1e-3)
    # Eka: penyebab = IPK (bukan kehadiran)
    kode = [f["kode"] for f in d["2024005"]["faktor"]]
    assert kode[0] == "GPA" and "attendance_rate" not in kode


def test_hitung_ulang_tidak_menggandakan_notifikasi(c, app, terisi):
    h = tok(c, "admin@x.id")
    c.post("/api/v1/risk/hitung", headers=h, json={"periode": "2026-1"})
    with app.app_context():
        assert Notifikasi.query.filter_by(kanal="wa").count() == 1
        assert RiskScore.query.count() == 4


def test_ringkasan_hitung_menyebut_high_tanpa_nomor(c, terisi):
    assert terisi["notifikasi_wa"] == 1
    assert terisi["high_tanpa_ortu"] == ["2024005"]  # orang tua Eka belum punya nomor WA
    assert terisi["high_tanpa_dosen_pa"] == []


def test_wa_dibuat_saat_nomor_ortu_diisi_kemudian(c, app, terisi):
    with app.app_context():
        OrangTua.query.filter_by(nim="2024005").first().no_wa = "6281200000055"
        db.session.commit()
    r = c.post("/api/v1/risk/hitung", headers=tok(c, "admin@x.id"), json={"periode": "2026-1"}).get_json()
    assert r["notifikasi_wa"] == 1 and r["high_tanpa_ortu"] == []


# ---------------------------------------------------------------- notifikasi WA
def test_pesan_wa_berisi_alasan_dan_kontak_dosen(c, terisi):
    h = tok(c, "dosen@x.id")
    daftar = c.get("/api/v1/notifikasi/wa-ortu", headers=h).get_json()
    assert daftar["provider"] == "manual" and len(daftar["notifikasi"]) == 1
    n = daftar["notifikasi"][0]
    assert n["no_tujuan"] == "6281200000003" and n["status"] == "pending"
    assert "RISIKO TINGGI" in n["isi"] and "Kehadiran rendah" in n["isi"] and "Citra Lestari" in n["isi"]
    assert "Dosen Satu" in n["isi"] and "6281200000001" in n["isi"]
    assert n["link"].startswith("https://wa.me/6281200000003?text=")
    assert "IPK" not in n["isi"]  # Citra: hanya karena kehadiran


def test_tandai_terkirim_dan_kepemilikan(c, terisi):
    nid = c.get("/api/v1/notifikasi/wa-ortu", headers=tok(c, "dosen@x.id")).get_json()["notifikasi"][0]["id"]
    # dosen lain tidak melihat / tidak bisa mengubah
    h2 = tok(c, "dosen2@x.id")
    assert c.get("/api/v1/notifikasi/wa-ortu", headers=h2).get_json()["notifikasi"] == []
    assert c.post(f"/api/v1/notifikasi/{nid}/terkirim", headers=h2).status_code == 404
    r = c.post(f"/api/v1/notifikasi/{nid}/terkirim", headers=tok(c, "dosen@x.id"))
    assert r.status_code == 200 and r.get_json()["status"] == "terkirim"
    assert c.post(f"/api/v1/notifikasi/{nid}/kirim-ulang", headers=tok(c, "dosen@x.id")).status_code == 409


def test_provider_log_langsung_terkirim(app):
    app.config["WA_PROVIDER"] = "log"
    c = app.test_client()
    h = tok(c, "admin@x.id")
    upload(c, h, "/api/v1/import/akademik", AKADEMIK)
    upload(c, h, "/api/v1/import/tugas", TUGAS)
    c.post("/api/v1/risk/hitung", headers=h, json={"periode": "2026-1"})
    n = c.get("/api/v1/notifikasi/wa-ortu", headers=h).get_json()["notifikasi"][0]
    assert n["status"] == "terkirim" and n["percobaan"] == 1


def test_provider_gagal_dicoba_ulang_lalu_gagal(app, monkeypatch):
    app.config.update(WA_PROVIDER="fonnte", FONNTE_TOKEN="x", WA_MAX_RETRY=3)
    import app.notif as notif

    def gagal(no, teks):
        raise RuntimeError("jaringan putus")
    monkeypatch.setattr(notif, "_kirim_fonnte", gagal)
    c = app.test_client()
    h = tok(c, "admin@x.id")
    upload(c, h, "/api/v1/import/akademik", AKADEMIK)
    upload(c, h, "/api/v1/import/tugas", TUGAS)
    c.post("/api/v1/risk/hitung", headers=h, json={"periode": "2026-1"})
    n = lambda: c.get("/api/v1/notifikasi/wa-ortu", headers=h).get_json()["notifikasi"][0]
    assert (n()["status"], n()["percobaan"]) == ("pending", 1)
    c.post("/api/v1/notifikasi/proses", headers=h)
    c.post("/api/v1/notifikasi/proses", headers=h)
    assert (n()["status"], n()["percobaan"]) == ("gagal", 3)
    # berhasil setelah kirim ulang
    monkeypatch.setattr(notif, "_kirim_fonnte", lambda no, teks: None)
    r = c.post(f"/api/v1/notifikasi/{n()['id']}/kirim-ulang", headers=h)
    assert r.get_json()["status"] == "terkirim"


def test_notifikasi_dashboard_dan_baca(c, terisi):
    hm = tok(c, "citra@x.id")
    isi = c.get("/api/v1/notifikasi", headers=hm).get_json()
    assert isi["belum_dibaca"] == 1 and "TINGGI" in isi["notifikasi"][0]["isi"]
    nid = isi["notifikasi"][0]["id"]
    assert c.patch(f"/api/v1/notifikasi/{nid}/baca", headers=hm).status_code == 200
    assert c.get("/api/v1/notifikasi", headers=hm).get_json()["belum_dibaca"] == 0
    # milik orang lain -> 404
    assert c.patch(f"/api/v1/notifikasi/{nid}/baca", headers=tok(c, "dosen@x.id")).status_code == 404


# ---------------------------------------------------------------- dashboard & RBAC
def test_dosen_hanya_melihat_bimbingannya(c, terisi):
    d1 = c.get("/api/v1/dashboard/dosen", headers=tok(c, "dosen@x.id")).get_json()
    nims1 = {m["nim"] for m in d1["mahasiswa"]}
    nims2 = {m["nim"] for m in c.get("/api/v1/dashboard/dosen", headers=tok(c, "dosen2@x.id")).get_json()["mahasiswa"]}
    assert nims1 == {"2024001", "2024003", "2024005"} and nims2 == {"2024009"}
    assert d1["mahasiswa"][0]["kategori"] == "High"  # urut: High dulu
    assert c.get("/api/v1/risk/2024009", headers=tok(c, "dosen@x.id")).status_code == 404


def test_mahasiswa_dan_ortu_hanya_data_sendiri(c, terisi):
    hm = tok(c, "citra@x.id")
    b = c.get("/api/v1/dashboard/mahasiswa", headers=hm).get_json()
    assert b["terbaru"]["kategori"] == "High" and b["dosen_pa"]["nama"] == "Dosen Satu"
    assert c.get("/api/v1/risk/2024001", headers=hm).status_code == 404
    assert c.get("/api/v1/dashboard/dosen", headers=hm).status_code == 403
    r = c.post("/api/v1/auth/login-ortu", json={"nim": "2024003", "email": "citra@x.id", "password": "rahasia123"})
    # ortu login memakai kredensial ortu sendiri; di sini pakai token ortu langsung
    ho = tok(c, "ortu-citra@x.id")
    o = c.get("/api/v1/dashboard/ortu", headers=ho).get_json()
    assert o["mahasiswa"]["nim"] == "2024003" and o["terbaru"]["faktor"][0]["kode"] == "attendance_rate"
    assert c.get("/api/v1/risk/2024005", headers=ho).status_code == 404
    assert c.get("/api/v1/dashboard/mahasiswa", headers=ho).status_code == 403


def test_kaprodi_agregat(c, terisi):
    b = c.get("/api/v1/dashboard/kaprodi", headers=tok(c, "kaprodi@x.id")).get_json()
    assert b["ringkasan"]["total"] == 4 and b["ringkasan"]["High"] == 2
    kode = {p["kode"]: p["jumlah"] for p in b["penyebab_high"]}
    assert kode["attendance_rate"] == 1 and kode["GPA"] == 1
    assert b["tren"][0]["periode"] == "2026-1"
    assert c.get("/api/v1/dashboard/kaprodi", headers=tok(c, "dosen@x.id")).status_code == 403


def test_risk_detail_riwayat(c, terisi):
    h = tok(c, "dosen@x.id")
    b = c.get("/api/v1/risk/2024003", headers=h).get_json()
    assert b["terbaru"]["kategori"] == "High" and len(b["riwayat"]) == 1


# ---------------------------------------------------------------- catatan
def test_catatan_dosen_dan_pesan_ke_mahasiswa(c, terisi):
    hd = tok(c, "dosen@x.id")
    r = c.post("/api/v1/catatan", headers=hd, json={"nim": "2024003", "isi_catatan": "Panggil untuk konseling",
                                                      "isi_pesan": "Temui saya Senin ya."})
    assert r.status_code == 201
    assert c.post("/api/v1/catatan", headers=tok(c, "dosen2@x.id"), json={"nim": "2024003", "isi_catatan": "xyz"}).status_code == 404
    hm = tok(c, "citra@x.id")
    lst = c.get("/api/v1/catatan", headers=hm).get_json()
    assert lst[0]["isi_pesan"] == "Temui saya Senin ya." and lst[0]["isi_catatan"] is None
    assert "Temui saya Senin" in c.get("/api/v1/notifikasi", headers=hm).get_json()["notifikasi"][0]["isi"]
    assert c.get("/api/v1/catatan?nim=2024003", headers=hd).get_json()[0]["isi_catatan"] == "Panggil untuk konseling"
    assert c.post("/api/v1/catatan", headers=hm, json={"nim": "2024003", "isi_catatan": "xyz"}).status_code == 403
    assert c.post("/api/v1/catatan", headers=hd, json={"nim": "2024003", "isi_catatan": ""}).status_code == 422


def test_catatan_tanpa_skor_ditolak(c):
    h = tok(c, "admin@x.id")
    upload(c, h, "/api/v1/import/akademik", AKADEMIK)
    r = c.post("/api/v1/catatan", headers=tok(c, "dosen@x.id"), json={"nim": "2024001", "isi_catatan": "catatan awal"})
    assert r.status_code == 409 and r.get_json()["code"] == "BELUM_ADA_SKOR"


# ---------------------------------------------------------------- laporan
def test_laporan_xlsx_pdf(c, terisi):
    h = tok(c, "kaprodi@x.id")
    x = c.get("/api/v1/laporan?periode=2026-1&format=xlsx", headers=h)
    assert x.status_code == 200 and x.data[:2] == b"PK"
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(x.data)).active
    assert ws.max_row == 5 and ws["A2"].value == "2024003"  # High lebih dulu
    p = c.get("/api/v1/laporan?periode=2026-1&format=pdf", headers=h)
    assert p.status_code == 200 and p.data[:4] == b"%PDF"
    # dosen hanya bimbingannya
    xd = c.get("/api/v1/laporan?format=xlsx", headers=tok(c, "dosen2@x.id"))
    assert load_workbook(io.BytesIO(xd.data)).active.max_row == 2
    assert c.get("/api/v1/laporan?format=xlsx", headers=tok(c, "citra@x.id")).status_code == 403
    assert c.get("/api/v1/laporan?format=doc", headers=h).status_code == 422


# ---------------------------------------------------------------- model langsung
def test_model_kasus_batas():
    from app import risk
    app = create_app({"TESTING": True, "SQLALCHEMY_DATABASE_URI": "sqlite:///:memory:"})
    with app.app_context():
        s, k, f = risk.prediksi({"GPA": 3.8, "avg_course_grade": 90, "attendance_rate": 0.97, "assignment_submission_rate": 0.98})
        assert k == "Low" and f == []
        s, k, f = risk.prediksi({"GPA": 3.0, "avg_course_grade": 76, "attendance_rate": 0.5, "assignment_submission_rate": 0.8})
        assert k == "High" and f[0]["kode"] == "attendance_rate"
        # data sebagian kosong tidak membuat error (diisi nilai tipikal)
        s, k, f = risk.prediksi({"GPA": None, "avg_course_grade": None, "attendance_rate": 0.5, "assignment_submission_rate": None})
        assert k == "High"
    assert risk.hitung_periode(__import__("datetime").date(2026, 9, 1)) == "2026-1"
    assert risk.hitung_periode(__import__("datetime").date(2027, 1, 15)) == "2026-1"
    assert risk.hitung_periode(__import__("datetime").date(2027, 3, 1)) == "2027-2"
